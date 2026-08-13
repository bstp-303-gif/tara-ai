import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

# Define required schema for each provider
SCHEMA = {
    'Google': ['Full Name', 'IC_Number', 'Email', 'School_Name', 'State', 'Certification', 'Cert_Level', 'Cert_Year'],
    'Microsoft': ['Teacher_Name', 'MyKad', 'Teacher_Email', 'Institution', 'State_Name', 'Cert_Programme', 'Cert_Type', 'Year_Achieved'],
    'Apple': ['Teacher_Full_Name', 'IC_Number', 'Email_Address', 'School', 'State', 'Apple_Programme', 'Level', 'Year']
}

NORMALIZE_MAPPING = {
    'Google': {
        'Full Name': 'name',
        'IC_Number': 'ic',
        'Email': 'email',
        'School_Name': 'school',
        'State': 'state',
        'Certification': 'certification',
        'Cert_Level': 'cert_level',
        'Cert_Year': 'cert_year'
    },
    'Microsoft': {
        'Teacher_Name': 'name',
        'MyKad': 'ic',
        'Teacher_Email': 'email',
        'Institution': 'school',
        'State_Name': 'state',
        'Cert_Programme': 'certification',
        'Cert_Type': 'cert_level',
        'Year_Achieved': 'cert_year'
    },
    'Apple': {
        'Teacher_Full_Name': 'name',
        'IC_Number': 'ic',
        'Email_Address': 'email',
        'School': 'school',
        'State': 'state',
        'Apple_Programme': 'certification',
        'Level': 'cert_level',
        'Year': 'cert_year'
    }
}


def validate_file(file_path, provider):
    """
    Validate Excel file structure and contents.
    Returns: {
        'is_valid': bool,
        'record_count': int,
        'error_message': str or None,
        'missing_values_map': dict,
        'dataframe': DataFrame or None
    }
    """
    try:
        df = pd.read_excel(file_path, header=0)

        required_columns = SCHEMA[provider]
        missing_columns = [col for col in required_columns if col not in df.columns]

        # Accept the file even if some columns are missing entirely — add
        # them as blank so the file can still be normalized/processed.
        # The gap is flagged (via error_message + missing_values_map) rather
        # than blocking the upload.
        for col in missing_columns:
            df[col] = None

        # Check for missing values in data (covers both blank cells and
        # columns that were missing outright)
        missing_values_map = {}
        for idx, row in df.iterrows():
            missing_cols = []
            for col in required_columns:
                cell_value = row[col]
                if pd.isna(cell_value) or (isinstance(cell_value, str) and cell_value.strip() == ''):
                    missing_cols.append(col)
            if missing_cols:
                missing_values_map[idx + 2] = missing_cols  # +2: header is row 1, data starts row 2

        warning_message = None
        if missing_columns:
            warning_message = f"Missing columns (accepted, left blank): {', '.join(missing_columns)}"

        return {
            'is_valid': True,
            'record_count': len(df),
            'error_message': warning_message,
            'missing_values_map': missing_values_map,
            'dataframe': df
        }

    except Exception as e:
        return {
            'is_valid': False,
            'record_count': 0,
            'error_message': f"Error reading file: {str(e)}",
            'missing_values_map': {},
            'dataframe': None
        }


def highlight_missing_values(input_file_path, output_file_path, missing_values_map):
    """
    Create a copy of the Excel file with missing values highlighted in red.
    """
    try:
        wb = load_workbook(input_file_path)
        ws = wb.active

        red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        white_font = Font(color="FFFFFF", bold=True)

        for row_num, missing_cols in missing_values_map.items():
            row_num = int(row_num)
            for col_name in missing_cols:
                for col_idx, cell in enumerate(ws[1], 1):
                    if cell.value == col_name:
                        cell_to_highlight = ws.cell(row=row_num, column=col_idx)
                        cell_to_highlight.fill = red_fill
                        cell_to_highlight.font = white_font
                        if cell_to_highlight.value is None:
                            cell_to_highlight.value = "[MISSING]"
                        break

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output_file_path)
        return True
    except Exception as e:
        print(f"Error highlighting file: {str(e)}")
        return False


def normalize_dataframe(df, provider):
    """
    Normalize dataframe columns to standard schema.
    """
    try:
        df_copy = df.copy()

        rename_map = NORMALIZE_MAPPING[provider]
        df_copy.rename(columns=rename_map, inplace=True)

        df_copy['provider'] = provider

        standard_cols = ['ic', 'name', 'email', 'school', 'state', 'certification', 'cert_level', 'cert_year', 'provider']
        df_copy = df_copy[[col for col in standard_cols if col in df_copy.columns]]

        return df_copy
    except Exception as e:
        print(f"Error normalizing dataframe: {str(e)}")
        return None


def save_normalized_file(dataframe, output_path):
    """
    Save normalized dataframe to Excel.
    """
    try:
        dataframe.to_excel(output_path, index=False, sheet_name='Normalized')

        wb = load_workbook(output_path)
        ws = wb.active

        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(output_path)
        return True
    except Exception as e:
        print(f"Error saving normalized file: {str(e)}")
        return False
